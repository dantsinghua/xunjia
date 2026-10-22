from io import BytesIO
from datetime import datetime
from zoneinfo import ZoneInfo
import json
from urllib.parse import parse_qs, urlsplit

from openpyxl import Workbook, load_workbook
from playwright.sync_api import expect


def test_quote_export_downloads_actual_xlsx_bytes_for_the_complete_filter(ui, tmp_path):
    page, url, state, requests, _ = ui
    workbook = Workbook()
    workbook.active.append(["code", "unique-value"])
    workbook.active.append(["000100.IB", "synthetic-source-value-2481"])
    output = BytesIO()
    workbook.save(output)
    state["export_bytes"] = output.getvalue()
    page.goto(url + "/#/quote")
    expect(page.locator(".mkt-row")).to_have_count(1)
    search = page.get_by_placeholder("债券代码 / 简称 / 发行人")
    search.fill("合成债券")
    search.press("Enter")
    with page.expect_download() as pending:
        page.get_by_test_id("quote-export").click()
    download = pending.value
    assert download.suggested_filename.endswith(".xlsx")
    target = tmp_path / "export.xlsx"
    download.save_as(target)
    assert target.read_bytes() == output.getvalue()
    book = load_workbook(target, read_only=True)
    assert book.active["B2"].value == "synthetic-source-value-2481"
    query = parse_qs(urlsplit([request.url for request in requests if "/api/bonds/export?" in request.url][-1]).query)
    chosen = json.loads(query["filter"][0])
    assert chosen["keyword"] == "合成债券"
    assert {"keyword", "chips", "prices", "channel", "broker", "groups"} <= chosen.keys()


def test_export_json_error_never_downloads_a_fake_xlsx(ui):
    page, url, _, _, _ = ui
    downloads = []
    page.on("download", lambda download: downloads.append(download))
    page.goto(url + "/#/quote")
    page.get_by_test_id("quote-export").click()
    expect(page.get_by_text("合成导出不可用", exact=True)).to_be_visible()
    assert downloads == []


def test_trade_board_displays_explicit_backend_facts_and_zero_gap(ui):
    page, url, state, _, _ = ui
    values = {"own_side": "SELL", "peer_quote": "2.10", "own_limit": "2.10", "gap_bp": "0",
              "peer_volume_wan": "1250.5", "quote_at": "2026-10-13T10:00:00+08:00",
              "freshness": "source-fresh", "phase": "询价中"}
    state["counterparties_data"] = [{"id": "qa-peer", "name": "合成对手甲", "bonds": [
        {"code": "000100.IB", "name": "合成债券", "stage": "询价中", "trade_day": datetime.now(ZoneInfo("Asia/Shanghai")).date().isoformat(), **values},
    ]}]
    page.goto(url + "/#/monitor")
    board = page.get_by_test_id("trade-board")
    expect(board).to_be_visible()
    expect(board.locator('[data-field="market_bid"], [data-field="market_ofr"]')).to_have_count(0)
    expect(board.locator('[data-field="own_side"]')).to_have_text("卖出 SELL")
    for name, value in values.items():
        if name != "own_side":
            expect(board.locator(f'[data-field="{name}"]')).to_have_text(value)


def test_trade_board_missing_fields_are_explicitly_unavailable(ui):
    page, url, _, _, _ = ui
    page.goto(url + "/#/monitor")
    board = page.get_by_test_id("trade-board")
    expect(board).to_be_visible()
    for field in ["own_side", "peer_quote", "own_limit", "gap_bp", "peer_volume_wan", "quote_at", "freshness", "phase"]:
        expect(board.locator(f'[data-field="{field}"]')).to_have_text("—")
