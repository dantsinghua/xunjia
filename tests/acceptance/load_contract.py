"""Explicit R32 load gate; never collected by the default test_*.py suite.

Run: pytest tests/acceptance/load_contract.py -q -s. Each backend owns a new
synthetic database; MySQL reuses only the audited new-container QA fixture.
"""
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timezone
import hashlib
from io import BytesIO
import json
import math
from pathlib import Path
import platform
import resource
import tempfile
import threading
import time

from fastapi.testclient import TestClient
from openpyxl import load_workbook
import pytest

from _db_observer import HistoryReadObserver
from mysql_contract import mysql_container, mysql_repo  # noqa: F401 - explicit fixture registration
from server.legacy_app import create_app
from server.auth import AuthConfig


NOW = 1_800_000_000
PER_OWNER = 100_000


def source_hashes():
    root = Path(__file__).resolve().parents[2]
    return {str(path.relative_to(root)): hashlib.sha256(path.read_bytes()).hexdigest()
            for path in sorted((root / "server").glob("*.py"))}


def seed(app):
    service = app.state.inquiry
    for index in range(5):
        owner, account = f"load-owner-{index}", f"99000{index}"
        service.repo.bind_account(owner, account, "synthetic-load-wb")
        draft = service.preview(owner, "synthetic-load-preview", [{"code": "000100.IB", "name": "合成负载债",
            "peer_ids": [f"load-peer-{peer}" for peer in range(10)], "our_side": "BUY", "volume_wan": "1000",
            "opening_yield_pct": "2.50", "limit_yield_pct": "2.40", "step_bp": "2"}])
        service.confirm(owner, draft["batch_id"], draft["version"], "synthetic-load-confirm")
        message = service.ingest(owner, {"event_id": "load-template", "account_uid": account, "peer_id": "load-peer-0",
            "source_seq": 1, "kind": "message", "text": "synthetic-template", "occurred_at": NOW, "historical": True})
        with service.repo.db.transaction() as tx:
            template = dict(tx.execute("SELECT * FROM xj_message WHERE message_id=?", (message["message_id"],)).fetchone())
            tx.execute("DELETE FROM xj_message WHERE message_id=?", (message["message_id"],))
            columns = list(template)
            placeholder = "%s" if service.repo.db.mysql else "?"
            sql = "INSERT INTO xj_message(" + ",".join(columns) + ") VALUES(" + ",".join(placeholder for _ in columns) + ")"
            cursor = tx.connection.cursor()
            try:
                for start in range(0, PER_OWNER, 5000):
                    rows = []
                    for number in range(start, min(start + 5000, PER_OWNER)):
                        row = dict(template, message_id=f"load-{index}-{number:06}", source_event_id=f"synthetic-{number}",
                                   source_seq=number + 1, text=f"synthetic-only-owner-{index}-message-{number:06}", occurred_at=NOW + number / 1000)
                        rows.append(tuple(row[key] for key in columns))
                    cursor.executemany(sql, rows)
            finally:
                cursor.close()


@pytest.mark.parametrize("backend", ["sqlite", "mysql"])
def test_five_owners_with_100000_history_rows_refresh_while_one_owner_exports(backend, tmp_path, request, monkeypatch):
    folder = Path(tempfile.mkdtemp(prefix="xunjia-load-acceptance-"))
    before = source_hashes()
    if backend == "mysql":
        _, database_url = request.getfixturevalue("mysql_repo")
    else:
        database_url = "sqlite:///" + str(folder / "synthetic-load.sqlite")
    report = {"backend": backend, "platform": platform.platform(), "owners": 5, "history_per_owner": PER_OWNER,
              "target_read_p95_seconds": 1.0, "scope": "isolated synthetic acceptance, not OceanBase capacity promise",
              "evidence_directory": str(folder), "source_before": before}
    (folder / "result.json").write_text(json.dumps(report, indent=2))
    print("Load evidence:", folder, flush=True)
    app = create_app(database_url=database_url, enable_background=False,
                     auth_config=AuthConfig(sso_secret="synthetic-load-sso", session_secret="synthetic-load-session-secret-minimum-32", clock=lambda: NOW))
    started = time.monotonic()
    seed(app)
    report["seed_seconds"] = round(time.monotonic() - started, 3)
    print("Seeded", backend, 5 * PER_OWNER, "synthetic history rows", flush=True)
    observer = HistoryReadObserver(monkeypatch, app.state.repository.db)
    readings, results = [], []
    lock, barrier = threading.Lock(), threading.Barrier(6)
    try:
        with TestClient(app) as client:
            headers = []
            for index in range(5):
                owner = f"load-owner-{index}"
                sign = hashlib.md5(f"{owner}|{owner}|{NOW}|synthetic-load-sso".encode()).hexdigest()
                response = client.post("/api/auth/sso", json={"userName": owner, "nickName": owner, "ts": str(NOW), "sign": sign})
                assert response.status_code == 200
                headers.append({"Authorization": "Bearer " + response.json()["token"]})

            def refresh(index):
                barrier.wait(timeout=20)
                for cycle in range(20):
                    for path in ("/api/monitor/summary", "/api/monitor/counterparties", "/api/monitor/chat?cp=load-peer-0&limit=50"):
                        begin = time.monotonic()
                        response = client.get(path, headers=headers[index])
                        elapsed = time.monotonic() - begin
                        assert response.status_code == 200, response.text
                        data = response.json()["data"]
                        if "/chat?" in path:
                            assert len(data["messages"]) == 50 and data["hasMore"] is True
                            assert all(message["text"].startswith(f"synthetic-only-owner-{index}-") for message in data["messages"])
                        elif path.endswith("counterparties"):
                            assert len(data) == 10
                        with lock:
                            readings.append({"owner": index, "path": path, "seconds": elapsed})

            def export():
                barrier.wait(timeout=20)
                begin = time.monotonic()
                response = client.get("/api/monitor/export", headers=headers[0], params={
                    "start": datetime.fromtimestamp(NOW, timezone.utc).isoformat(),
                    "end": datetime.fromtimestamp(NOW + 101, timezone.utc).isoformat(), "peer": "load-peer-0"})
                assert response.status_code == 200, response.text[:300]
                results.append({"export_seconds": time.monotonic() - begin, "xlsx_bytes": len(response.content)})
                (folder / "synthetic-export.xlsx").write_bytes(response.content)
                return response.content

            with ThreadPoolExecutor(max_workers=6) as pool:
                workers = [pool.submit(refresh, index) for index in range(5)]
                exported = pool.submit(export)
                for worker in workers:
                    worker.result(timeout=600)
                workbook = load_workbook(BytesIO(exported.result(timeout=600)), read_only=True)
                count = 0
                for row in workbook.active.iter_rows(min_row=2):
                    assert row[2].value.startswith("synthetic-only-owner-0-")
                    assert row[2].data_type == "s"
                    count += 1
                workbook.close()
                assert count == PER_OWNER
            report.update(results[0], export_rows=count)
        ordered = sorted(row["seconds"] for row in readings)
        p95 = ordered[math.ceil(len(ordered) * 0.95) - 1]
        report.update(read_requests=len(readings), read_p95_seconds=p95, read_max_seconds=max(ordered),
                      read_latencies=readings, history_cursor_reads=observer.reads,
                      max_rss_kib=resource.getrusage(resource.RUSAGE_SELF).ru_maxrss)
        assert len(readings) == 300
        chunks = [row for row in observer.reads if row["method"] == "fetchmany"]
        assert sum(row["count"] for row in chunks) == PER_OWNER
        assert max(row["count"] for row in chunks) <= 500
        assert p95 <= 1.0, f"{backend} read P95 {p95:.3f}s exceeds the design acceptance target of 1s"
        report["passed"] = True
    finally:
        report["source_after"] = source_hashes()
        report["source_drift"] = sorted(path for path in before.keys() | report["source_after"].keys() if before.get(path) != report["source_after"].get(path))
        report.setdefault("passed", False)
        (folder / "result.json").write_text(json.dumps(report, indent=2))
        print(json.dumps({key: report.get(key) for key in ("backend", "passed", "read_requests", "read_p95_seconds", "export_rows", "export_seconds", "source_drift")}), flush=True)
