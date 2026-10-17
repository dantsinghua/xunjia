"""Authenticated public market, owner overlays, exact export and websocket boundary."""
import hashlib
from io import BytesIO
import json

from fastapi.testclient import TestClient
from openpyxl import load_workbook
import pytest
from starlette.websockets import WebSocketDisconnect

from server.legacy_app import create_app
from server.auth import AuthConfig
from server.market import MarketService
from server.repository import Repository


@pytest.fixture
def market_http(tmp_path):
    url = "sqlite:///" + str(tmp_path / "market-business.db")
    repo = Repository(url)
    for user, uid in [("a", "910001"), ("b", "910002")]:
        repo.bind_account(user, uid, "synthetic-wb")
    rows = [{"code": f"{number:06d}.IB", "name": "合成券-" + str(number), "bondUniCode": str(number + 10000),
             "flowId": number, "ob": "2.10", "oo": "2.20", "tags": [], "clearing": "+1",
             "ib": "PRIVATE-INTENT", "status": "已成交"} for number in range(1, 2052)]
    clock = [1_800_000_000]
    app = create_app(database_url=url, market=MarketService(repo, rows_provider=lambda channel, broker: rows),
                     auth_config=AuthConfig(sso_secret="sso", session_secret="market-test-session-secret-over-32", token_ttl_seconds=30, clock=lambda: clock[0]))
    with TestClient(app) as client:
        tokens = {}
        for user in ("a", "b"):
            sign = hashlib.md5(f"{user}|{user}|{clock[0]}|sso".encode()).hexdigest()
            tokens[user] = client.post("/api/auth/sso", json={"userName": user, "nickName": user, "ts": str(clock[0]), "sign": sign}).json()["token"]
        yield app, client, tokens, rows, clock


def headers(token):
    return {"Authorization": "Bearer " + token}


@pytest.fixture
def sql_market_http(market_http, tmp_path, monkeypatch):
    from server import market_cache
    from server.quote_runtime import QBMarket

    app, client, tokens, _, clock = market_http
    monkeypatch.setattr(market_cache, "_conn", None)
    cache_path = str(tmp_path / "independent-public-cache.db")
    static = [{"bond_uni_code": str(i), "windcode": f"{i:06d}.IB", "sec_name": "SQL合成券" + str(i),
               "bond_type": "中票", "residual_days": 730, "tags": ["绿色债"] if i % 2 == 0 else [],
               "prov": "广东" if i % 3 == 0 else "江苏", "val_yield": 2.0} for i in range(1, 138)]
    connection = market_cache.init_db(static, path=cache_path)
    provider = QBMarket(cache_path)
    market_cache.upsert_quotes([{"bond_uni_code": str(i), "group_name": "bond-all", "issue_time": clock[0],
                                "best_bid": 2.0 + (i % 5) / 100, "best_ofr": 1.90, "best_bid_vol": 1000,
                                "best_ofr_vol": 2000, "firmness": "实盘", "clearing": "+1",
                                "ib": 99.9, "status": "已成交"} for i in range(1, 138)])
    # Equal primary cursor values exercise the necessary secondary key.
    with market_cache.transaction() as tx:
        tx.execute("UPDATE quote_latest SET flow_id=7")
    app.state.market.rows_provider = provider.rows
    try:
        yield app, client, tokens, connection
    finally:
        connection.close()


def test_sql_backed_http_pages_visit_equal_version_rows_exactly_once(sql_market_http):
    _, client, tokens, _ = sql_market_http
    cursor, seen = {}, []
    for _ in range(30):
        response = client.get("/api/bonds", headers=headers(tokens["a"]), params={"limit": 7, **cursor})
        assert response.status_code == 200
        body = response.json()
        assert body["total"] == 137
        assert len(body["data"]) <= 7
        seen.extend(row["code"] for row in body["data"])
        if not body["hasMore"]:
            break
        cursor = {"before": body["nextBefore"], "beforeUni": body["nextUni"]}
    assert len(seen) == len(set(seen)) == 137
    assert set(seen) == {f"{i:06d}.IB" for i in range(1, 138)}


def test_sql_filter_pagination_selection_and_export_agree_on_combined_constraints(sql_market_http):
    _, client, tokens, _ = sql_market_http
    selected = {"keyword": "SQL合成", "chips": {"标签": 1, "清算": 2, "期限": 2, "报价类型": 1},
                "prices": {"Bid": {"lo": "2.01", "hi": "2.02"}, "Ofr": {"lo": "1.9", "hi": "1.9"},
                           "Bid-中债": {"lo": "1", "hi": "2"}, "中债-Ofr": {"lo": "10", "hi": "10"},
                           "买量": {"lo": "1000", "hi": "1000"}, "卖量": {"lo": "2000", "hi": "2000"}},
                "channel": "best", "broker": "", "groups": ["bond-prov_广东", "bond-type_中票"]}
    expected = {f"{i:06d}.IB" for i in range(1, 138) if i % 6 == 0 and i % 5 in (1, 2)}
    seen, cursor = [], {}
    for _ in range(20):
        response = client.get("/api/bonds", headers=headers(tokens["a"]), params={"filter": json.dumps(selected), "limit": 3, **cursor})
        assert response.status_code == 200
        body = response.json()
        assert body["total"] == len(expected)
        seen.extend(row["code"] for row in body["data"])
        if not body["hasMore"]:
            break
        cursor = {"before": body["nextBefore"], "beforeUni": body["nextUni"]}
    assert set(seen) == expected and len(seen) == len(expected)
    omitted = sorted(expected)[0]
    resolved = client.post("/api/bonds/resolve", headers=headers(tokens["a"]), json={"selection": {"selectAll": True, "filter": selected, "except": [omitted]}})
    assert resolved.status_code == 200
    assert {row["code"] for row in resolved.json()["data"]} == expected - {omitted}
    exported = client.get("/api/bonds/export", headers=headers(tokens["a"]), params={"filter": json.dumps(selected)})
    assert exported.status_code == 200
    book = load_workbook(BytesIO(exported.content), read_only=True)
    assert {row[0].value for row in list(book.active.rows)[1:]} == expected


def test_sql_status_filter_is_derived_only_from_authenticated_owners_sessions(sql_market_http):
    from server.business import InquiryService

    app, client, tokens, _ = sql_market_http
    service = InquiryService(app.state.repository)
    draft = service.preview("a", "sql-owner-preview", [{"code": "000001.IB", "name": "SQL合成券1", "peer_ids": ["sql-peer"],
                            "our_side": "BUY", "volume_wan": "1000", "opening_yield_pct": "2.50", "limit_yield_pct": "2.40", "step_bp": "2"}])
    service.confirm("a", draft["batch_id"], draft["version"], "sql-owner-confirm")
    pending = json.dumps({"chips": {"当前状态": 3}})
    own = client.get("/api/bonds", headers=headers(tokens["a"]), params={"filter": pending}).json()
    foreign = client.get("/api/bonds", headers=headers(tokens["b"]), params={"filter": pending, "user_id": "a"}).json()
    assert own["total"] == 1 and [row["code"] for row in own["data"]] == ["000001.IB"]
    assert own["data"][0]["ib"] is None
    assert foreign["total"] == 0 and foreign["data"] == []
    unfiltered = client.get("/api/bonds", headers=headers(tokens["b"])).json()
    assert {row["status"] for row in unfiltered["data"]} == {"待询价"}


def test_readonly_question_uses_real_market_values_and_omits_missing_chart_points(market_http):
    _, client, tokens, rows, _ = market_http
    rows[0]["ob"], rows[0]["oo"] = None, 0
    response = client.post("/api/qa", headers=headers(tokens["a"]), json={"question": "报价 000001.IB"})
    assert response.status_code == 200
    data = response.json()["data"]
    assert "Bid：未接入或缺失" in data["answer"]
    assert "Ofr：0%" in data["answer"]
    assert data["chart"]["categories"] == ["Ofr"]
    assert data["chart"]["series"][0]["data"] == [0]
    rows[0]["oo"] = None
    response = client.post("/api/qa", headers=headers(tokens["a"]), json={"question": "报价 000001.IB"})
    assert response.status_code == 200 and "chart" not in response.json()["data"]


def test_readonly_search_has_actual_candidates_but_never_fabricates_unavailable_reference_facts(market_http):
    _, client, tokens, _, _ = market_http
    response = client.post("/api/qa", headers=headers(tokens["a"]), json={"question": "查券 000001.IB"})
    assert response.status_code == 200
    answer = response.json()["data"]["answer"]
    assert "000001.IB" in answer and "合成券-1" in answer
    assert "参考资料尚未接入或当前不可用" in answer
    assert "chart" not in response.json()["data"]


@pytest.mark.parametrize("body", [{"question": "SELECT * FROM xj_binding"}, {"question": "我的在谈概要", "user_id": "a"}, {"question": []}, {"question": "x" * 201}])
def test_readonly_question_rejects_sql_identity_override_and_invalid_shapes(market_http, body):
    _, client, tokens, _, _ = market_http
    response = client.post("/api/qa", headers=headers(tokens["b"]), json=body)
    assert response.status_code in (400, 422)


def test_readonly_summary_only_reads_own_live_sessions_without_creating_commands(market_http):
    from server.business import InquiryService

    app, client, tokens, _, _ = market_http
    service = InquiryService(app.state.repository)
    for owner, code in (("a", "000001.IB"), ("b", "000002.IB")):
        draft = service.preview(owner, "qa-preview", [{"code": code, "name": "合成问答券", "peer_ids": ["private-peer-" + owner],
                "our_side": "BUY", "volume_wan": "1000", "opening_yield_pct": "2.50", "limit_yield_pct": "2.40", "step_bp": "2"}])
        service.confirm(owner, draft["batch_id"], draft["version"], "qa-confirm")
    before = service.commands("a"), service.commands("b")
    response = client.post("/api/qa", headers=headers(tokens["a"]), json={"question": "我的在谈概要"})
    assert response.status_code == 200
    assert "000001.IB" in response.text and "private-peer-a" in response.text
    assert "000002.IB" not in response.text and "private-peer-b" not in response.text
    unsupported = client.post("/api/qa", headers=headers(tokens["a"]), json={"question": "立即成交所有券"})
    assert unsupported.status_code == 200 and "暂不支持" in unsupported.text
    assert "chart" not in unsupported.json()["data"]
    assert (service.commands("a"), service.commands("b")) == before
    assert client.post("/api/qa", json={"question": "我的在谈概要"}).status_code == 401


def test_all_2051_rows_are_reachable_without_duplicate_or_skipped_pages(market_http):
    _, client, tokens, _, _ = market_http
    before, before_uni, seen = 0, "", []
    for _ in range(50):
        response = client.get("/api/bonds", headers=headers(tokens["a"]), params={"limit": 50, "before": before, "beforeUni": before_uni})
        assert response.status_code == 200
        body = response.json()
        assert body["total"] == 2051
        seen.extend(row["code"] for row in body["data"])
        if not body["hasMore"]:
            break
        before, before_uni = body["nextBefore"], body["nextUni"]
    assert len(seen) == len(set(seen)) == 2051
    assert "000001.IB" in seen and "002051.IB" in seen


def test_filter_list_resolve_and_export_use_the_same_actual_set(market_http):
    _, client, tokens, rows, _ = market_http
    rows[20]["name"] = "独立目标甲"
    rows[2000]["name"] = "独立目标乙"
    selected = {"keyword": "独立目标", "chips": {"清算": 2}, "prices": {"Bid": {"lo": "2.00", "hi": "2.15"}},
                "channel": "best", "broker": "", "groups": ["bond-all"]}
    listed = client.get("/api/bonds", headers=headers(tokens["a"]), params={"filter": json.dumps(selected)}).json()["data"]
    resolved = client.post("/api/bonds/resolve", headers=headers(tokens["a"]), json={"selection": {"selectAll": True, "filter": selected, "except": []}})
    assert resolved.status_code == 200
    expected = {row["code"] for row in listed}
    assert expected == {"000021.IB", "002001.IB"}
    assert {row["code"] for row in resolved.json()["data"]} == expected
    exported = client.get("/api/bonds/export", headers=headers(tokens["a"]), params={"filter": json.dumps(selected)})
    assert exported.status_code == 200
    workbook = load_workbook(BytesIO(exported.content), read_only=True)
    assert {row[0].value for row in list(workbook.active.rows)[1:]} == expected


def test_excluding_a_market_code_never_uses_internal_unicode_id(market_http):
    _, client, tokens, rows, _ = market_http
    rows[10]["name"] = rows[20]["name"] = "两条目标"
    selected = {"selectAll": True, "filter": {"keyword": "两条目标"}, "except": [rows[10]["code"]]}
    result = client.post("/api/bonds/resolve", headers=headers(tokens["a"]), json={"selection": selected})
    assert result.status_code == 200
    assert [row["code"] for row in result.json()["data"]] == [rows[20]["code"]]


def test_public_provider_cannot_inject_other_traders_intents_or_status(market_http):
    _, client, tokens, _, _ = market_http
    for user in ("a", "b"):
        response = client.get("/api/bonds", headers=headers(tokens[user]))
        assert "PRIVATE-INTENT" not in response.text
        assert {row["status"] for row in response.json()["data"]} == {"待询价"}


def test_exported_formula_like_name_is_a_literal_cell(market_http):
    _, client, tokens, rows, _ = market_http
    rows[0]["name"] = '=HYPERLINK("https://invalid.example","synthetic")'
    response = client.get("/api/bonds/export", headers=headers(tokens["a"]), params={"filter": json.dumps({"keyword": "HYPERLINK"})})
    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.content), data_only=False)
    cell = workbook.active["B2"]
    assert cell.value == rows[0]["name"]
    assert cell.data_type == "s"


def test_websocket_requires_authentication_before_subscription(market_http):
    _, client, _, _, _ = market_http
    with client.websocket_connect("/ws/quote") as websocket:
        websocket.send_json({"op": "sub", "groups": ["bond-all"]})
        with pytest.raises(WebSocketDisconnect) as error:
            websocket.receive_json()
        assert error.value.code == 4401


def test_websocket_rejects_disallowed_cross_origin(market_http):
    _, client, _, _, _ = market_http
    with pytest.raises(WebSocketDisconnect) as error:
        with client.websocket_connect("/ws/quote", headers={"Origin": "https://unauthorized.invalid"}):
            pytest.fail("An unauthorized origin must not establish the connection")
    assert error.value.code == 4403


def test_websocket_expiry_closes_an_existing_subscription(market_http):
    _, client, tokens, _, clock = market_http
    with client.websocket_connect("/ws/quote") as websocket:
        websocket.send_json({"op": "auth", "token": tokens["a"]})
        assert websocket.receive_json() == {"op": "auth", "ok": True}
        clock[0] += 31
        with pytest.raises(WebSocketDisconnect) as error:
            websocket.receive_json()
        assert error.value.code == 4401


def test_websocket_public_projection_excludes_private_fields_even_inside_static_data(market_http):
    app, client, tokens, _, _ = market_http
    with client.websocket_connect("/ws/quote") as websocket:
        websocket.send_json({"op": "auth", "token": tokens["a"]})
        assert websocket.receive_json()["ok"] is True
        websocket.send_json({"op": "sub", "groups": ["bond-all"]})
        websocket.send_json({"op": "ping"})
        assert websocket.receive_json()["op"] == "pong"
        app.state.quote_hub.publish({"f": 1, "u": "10001", "t": 123, "bb": None, "bo": 2.20,
                                     "ib": "SECRET-INTENT", "s": {"bond_uni_code": "10001", "windcode": "000001.IB",
                                                                       "sec_name": "合成券", "internal_credit": "SECRET-CREDIT"}}, {"bond-all"})
        result = websocket.receive_json()
        assert result["op"] == "batch"
        assert result["items"][0]["bb"] is None
        assert "SECRET-" not in json.dumps(result)


@pytest.mark.parametrize("path", ["/api/bond-static/all", "/api/bond-static?codes=10001"])
def test_public_static_http_projection_excludes_internal_fields(market_http, path):
    app, client, tokens, _, _ = market_http
    app.state.market.static_provider = lambda: [{"bond_uni_code": "10001", "windcode": "000001.IB", "sec_name": "合成券",
                                               "internal_credit": "SECRET-CREDIT", "raw_json": {"password": "SECRET-PASSWORD"}}]
    response = client.get(path, headers=headers(tokens["a"]))
    assert response.status_code == 200
    assert response.json()["data"][0]["windcode"] == "000001.IB"
    assert "SECRET-" not in response.text
