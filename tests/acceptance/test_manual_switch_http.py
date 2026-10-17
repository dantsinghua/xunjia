"""Manual authorization and explicit channel switching use production routes."""
import hashlib
from datetime import datetime
from io import BytesIO

from fastapi.testclient import TestClient
import pytest
from openpyxl import load_workbook

from relay.protocol import Binding
from server.legacy_app import create_app
from server.auth import AuthConfig
from server.runtime import Dispatcher
from server.transports import TransportRouter


NOW = 1_800_000_000


class Adapter:
    def __init__(self, binding):
        self.binding, self.sent, self.result = binding, [], "sent"

    def catalog(self):
        return {"account_uid": self.binding.account_uid, "workbench_id": self.binding.workbench_id}

    def contacts(self):
        return {"account_uid": self.binding.account_uid, "contacts": [], "org": {}}

    def send(self, command):
        self.sent.append(command.copy())
        return {key: command[key] for key in ("command_id", "account_uid", "peer_id", "content_hash")} | {"status": self.result, "evidence": "synthetic-evidence"}

    def reconcile(self, _command):
        return None

    def cancel(self, command):
        return {key: command[key] for key in ("command_id", "account_uid", "peer_id", "content_hash")} | {"status": "unknown", "evidence": "cancellation-not-yet-proven"}


@pytest.fixture
def manual_app(tmp_path):
    adapters = {}
    for owner, account in [("a", "910001"), ("b", "910002")]:
        binding = Binding(owner, account, "synthetic-wb", b"0123456789abcdef", b"s" * 32)
        for route in ("oss", "mail"):
            adapters[(owner, account, route)] = Adapter(binding)
    transport = TransportRouter(adapters)
    app = create_app(database_url="sqlite:///" + str(tmp_path / "manual.db"), enable_background=False, transport=transport,
                     auth_config=AuthConfig(sso_secret="sso", session_secret="synthetic-session-secret-over-32", clock=lambda: NOW))
    with TestClient(app) as client:
        headers = {}
        for owner, account in [("a", "910001"), ("b", "910002")]:
            app.state.repository.bind_account(owner, account, "synthetic-wb")
            sign = hashlib.md5(f"{owner}|{owner}|{NOW}|sso".encode()).hexdigest()
            result = client.post("/api/auth/sso", json={"userName": owner, "nickName": owner, "ts": str(NOW), "sign": sign})
            headers[owner] = {"Authorization": "Bearer " + result.json()["token"]}
        yield app, client, headers, adapters, transport


def authorize(case):
    service = case[0].state.inquiry
    draft = service.preview("a", "first", [{"code": "000100.IB", "name": "合成债", "peer_ids": ["peer"],
                           "our_side": "BUY", "volume_wan": "1000", "opening_yield_pct": "2.50", "limit_yield_pct": "2.40", "step_bp": "2"}])
    return service.confirm("a", draft["batch_id"], draft["version"], "confirm")["commands"][0]


def test_archive_is_owner_scoped_idempotent_and_preserves_other_sessions_and_messages(manual_app):
    app, client, headers, _, _ = manual_app
    authorize(manual_app)
    service = app.state.inquiry
    draft = service.preview("a", "second", [{"code": "000200.IB", "name": "合成第二券", "peer_ids": ["peer"],
                            "our_side": "BUY", "volume_wan": "1000", "opening_yield_pct": "2.50", "limit_yield_pct": "2.40", "step_bp": "2"}])
    service.confirm("a", draft["batch_id"], draft["version"], "second-confirm")
    sessions = service.sessions("a")
    target = next(row for row in sessions if row["bond_code"] == "000100.IB")["session_id"]
    service.ingest("a", {"event_id": "archive-history", "account_uid": "910001", "peer_id": "peer", "source_seq": 1,
                         "kind": "message", "text": "归档后仍保留的历史", "occurred_at": NOW, "historical": True, "direction": "in"})
    assert client.post(f"/api/sessions/{target}/archive", headers=headers["b"]).status_code == 403
    assert client.post(f"/api/sessions/{target}/archive", headers=headers["a"]).status_code == 400
    service.set_bot("a", False)
    commands = service.commands("a")
    for _ in range(2):
        response = client.post(f"/api/sessions/{target}/archive", headers=headers["a"])
        assert response.status_code == 200 and response.json()["data"]["archived"] == target
    assert [row["bond_code"] for row in service.sessions("a")] == ["000200.IB"]
    assert service.commands("a") == commands
    history = client.get("/api/monitor/chat", headers=headers["a"], params={"cp": "peer"})
    assert "归档后仍保留的历史" in history.text
    assert client.delete("/api/im/inquiry/000200.IB", headers=headers["a"]).status_code in (404, 405)


@pytest.mark.parametrize("status", ["pending", "unknown"])
def test_stopped_session_with_unresolved_execution_cannot_be_archived(manual_app, status):
    app, client, headers, adapters, transport = manual_app
    authorize(manual_app)
    adapters[("a", "910001", "oss")].result = status
    Dispatcher(app.state.inquiry, transport, clock=lambda: NOW).run_once()
    app.state.inquiry.set_bot("a", False)
    session = app.state.inquiry.sessions("a")[0]
    assert session["state"] == "STOPPED"
    response = client.post(f"/api/sessions/{session['session_id']}/archive", headers=headers["a"])
    assert response.status_code == 400
    assert len(app.state.inquiry.sessions("a")) == 1


def test_explicit_manual_send_takes_control_and_can_send_when_bot_is_stopped(manual_app):
    app, client, headers, adapters, transport = manual_app
    initial = authorize(manual_app)
    app.state.inquiry.set_bot("a", False)
    response = client.post("/api/monitor/send", headers=headers["a"], json={"cp": "peer", "text": "人工明确询问，请稍后回复", "request_id": "manual-one"})
    assert response.status_code == 200, response.text
    assert response.json()["data"]["status"] == "queued"
    assert app.state.inquiry.conversations("a")[0]["mode"] == "HUMAN"
    Dispatcher(app.state.inquiry, transport, clock=lambda: NOW).run_once()
    sent = adapters[("a", "910001", "oss")].sent
    assert len(sent) == 1 and sent[0]["kind"] == "manual_send"
    assert sent[0]["command_id"] != initial["command_id"]
    assert app.state.repository.binding_for("a")["bot_enabled"] == 0


def test_manual_takeover_cancels_queued_bot_work_atomically(manual_app):
    app, client, headers, _, _ = manual_app
    initial = authorize(manual_app)
    response = client.post("/api/monitor/send", headers=headers["a"], json={"cp": "peer", "text": "人工接着沟通", "request_id": "manual-one"})
    assert response.status_code == 200
    commands = {row["command_id"]: row for row in app.state.inquiry.commands("a")}
    assert commands[initial["command_id"]]["status"] == "cancelled"
    assert commands[response.json()["data"]["cmd_id"]]["status"] == "queued"
    assert app.state.inquiry.conversations("a")[0]["mode"] == "HUMAN"


@pytest.mark.parametrize("text", ["这笔确认成交 done", "我们的内部底价是2.40%"])
def test_manual_path_cannot_bypass_trade_or_secret_restrictions(manual_app, text):
    app, client, headers, _, _ = manual_app
    authorize(manual_app)
    before = app.state.inquiry.conversations("a")
    response = client.post("/api/monitor/send", headers=headers["a"], json={"cp": "peer", "text": text, "request_id": "bad-manual"})
    assert response.status_code == 400
    assert len(app.state.inquiry.commands("a")) == 1
    assert app.state.inquiry.conversations("a") == before


def test_manual_retry_preserves_id_and_changed_text_conflicts(manual_app):
    app, client, headers, _, _ = manual_app
    authorize(manual_app)
    body = {"cp": "peer", "text": "请稍后回复", "request_id": "manual-stable"}
    one = client.post("/api/monitor/send", headers=headers["a"], json=body)
    two = client.post("/api/monitor/send", headers=headers["a"], json=body)
    assert one.status_code == two.status_code == 200
    assert one.json()["data"]["cmd_id"] == two.json()["data"]["cmd_id"]
    assert client.post("/api/monitor/send", headers=headers["a"], json={**body, "text": "变更内容"}).status_code in (400, 409)
    assert len(app.state.inquiry.commands("a")) == 2


def test_foreign_owner_cannot_manually_send_on_another_owners_conversation(manual_app):
    app, client, headers, _, _ = manual_app
    authorize(manual_app)
    response = client.post("/api/monitor/send", headers=headers["b"], json={"cp": "peer", "text": "越权尝试", "request_id": "foreign", "owner_user_id": "a"})
    assert response.status_code in (400, 403, 404)
    assert app.state.inquiry.commands("b") == []
    assert len(app.state.inquiry.commands("a")) == 1


def test_unknown_result_blocks_manual_bypass_and_channel_switch_even_after_stop(manual_app):
    app, client, headers, adapters, transport = manual_app
    authorize(manual_app)
    adapters[("a", "910001", "oss")].result = "unknown"
    Dispatcher(app.state.inquiry, transport, clock=lambda: NOW).run_once()
    app.state.inquiry.set_bot("a", False)
    version = app.state.repository.binding_for("a")["version"]
    manual = client.post("/api/monitor/send", headers=headers["a"], json={"cp": "peer", "text": "不应绕过未决任务", "request_id": "manual-unknown"})
    switch = client.post("/api/im/links/910001:mail/activate", headers=headers["a"], json={"expected_version": version})
    assert manual.status_code == switch.status_code == 400
    assert app.state.repository.binding_for("a")["transport"] == "oss"
    assert adapters[("a", "910001", "mail")].sent == []


def test_idle_stopped_binding_can_switch_explicitly_but_stale_version_cannot(manual_app):
    app, client, headers, _, _ = manual_app
    version = app.state.repository.binding_for("a")["version"]
    response = client.post("/api/im/links/910001:mail/activate", headers=headers["a"], json={"expected_version": version})
    assert response.status_code == 200
    assert app.state.repository.binding_for("a")["transport"] == "mail"
    stale = client.post("/api/im/links/910001:oss/activate", headers=headers["a"], json={"expected_version": version})
    assert stale.status_code == 400
    assert app.state.repository.binding_for("a")["transport"] == "mail"


def test_foreign_link_id_cannot_switch_another_owners_channel(manual_app):
    app, client, headers, _, _ = manual_app
    response = client.post("/api/im/links/910002:mail/activate", headers=headers["a"], json={"expected_version": 1})
    assert response.status_code == 403
    assert app.state.repository.binding_for("b")["transport"] == "oss"


def test_logged_in_trader_cannot_create_a_different_account_binding(manual_app):
    app, client, headers, _, _ = manual_app
    response = client.post("/api/im/links", headers=headers["a"], json={"account_uid": "stolen-account", "workbench_id": "other", "mode": "oss"})
    assert response.status_code in (403, 404, 405)
    assert app.state.repository.binding_for("a")["account_uid"] == "910001"


def test_chat_export_honors_explicit_timezone_and_excludes_other_owners(manual_app):
    app, client, headers, _, _ = manual_app
    at = datetime.fromisoformat("2026-10-13T09:00:00+08:00").timestamp()
    for owner, uid, text in [("a", "910001", "本人时区范围内消息"), ("b", "910002", "他人私聊不可导出")]:
        app.state.inquiry.ingest(owner, {"event_id": "export-event", "account_uid": uid, "peer_id": "peer",
                                       "kind": "private", "source_seq": 1, "text": text, "occurred_at": at, "historical": True})
    response = client.get("/api/monitor/export", headers=headers["a"], params={
        "start": "2026-10-13T08:30:00+08:00", "end": "2026-10-13T09:30:00+08:00",
    })
    assert response.status_code == 200
    book = load_workbook(BytesIO(response.content), read_only=True)
    texts = [row[2].value for row in list(book.active.rows)[1:]]
    assert texts == ["本人时区范围内消息"]


@pytest.mark.parametrize("params", [{}, {"start": "2026-10-13"}, {"end": "2026-10-14"},
                                      {"start": "not-date", "end": "2026-10-14"},
                                      {"start": "2026-10-15", "end": "2026-10-13"},
                                      {"start": "2026-10-13", "end": "2026-10-14", "peer": "x" * 161}])
def test_history_export_requires_valid_explicit_scope(manual_app, params):
    _, client, headers, _, _ = manual_app
    response = client.get("/api/monitor/export", headers=headers["a"], params=params)
    assert response.status_code in (400, 422)
    assert "spreadsheet" not in response.headers.get("content-type", "")


def test_history_export_reads_500_row_chunks_filters_peer_and_never_writes_formula_cells(manual_app, monkeypatch):
    from _db_observer import HistoryReadObserver

    app, client, headers, _, _ = manual_app
    at = datetime.fromisoformat("2026-10-13T09:00:00+08:00").timestamp()
    for owner, uid, peer in [("a", "910001", "included-peer"), ("a", "910001", "excluded-peer"), ("b", "910002", "included-peer")]:
        app.state.inquiry.ingest(owner, {"event_id": "seed-" + peer, "account_uid": uid, "peer_id": peer,
                                       "kind": "message", "source_seq": 1, "text": "seed", "occurred_at": at, "historical": True})
    with app.state.repository.db.transaction() as tx:
        source = dict(tx.execute("SELECT * FROM xj_message WHERE owner_user_id=? AND source_event_id=?", ("a", "seed-included-peer")).fetchone())
        tx.execute("DELETE FROM xj_message WHERE message_id=?", (source["message_id"],))
        columns = list(source)
        sql = "INSERT INTO xj_message(" + ",".join(columns) + ") VALUES(" + ",".join("?" for _ in columns) + ")"
        for i in range(1001):
            row = dict(source, message_id=f"export-{i:04}", source_event_id=f"export-event-{i}", source_seq=i + 1,
                       text='=HYPERLINK("https://synthetic.invalid","literal")' if i == 500 else f"included-{i}")
            tx.execute(sql, tuple(row[key] for key in columns))
    observer = HistoryReadObserver(monkeypatch, app.state.repository.db)
    response = client.get("/api/monitor/export", headers=headers["a"], params={"start": "2026-10-13", "end": "2026-10-13", "peer": "included-peer"})
    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.content), read_only=True)
    rows = list(workbook.active.rows)[1:]
    assert len(rows) == 1001
    formula = [row[2] for row in rows if str(row[2].value).startswith("=HYPERLINK")]
    assert len(formula) == 1 and formula[0].data_type == "s"
    assert all(row[2].value != "seed" for row in rows)
    assert [read["count"] for read in observer.reads if read["method"] == "fetchmany"] == [500, 500, 1, 0]


@pytest.mark.parametrize("boundary", ["rows", "characters"])
def test_export_hard_limits_reject_instead_of_silently_returning_a_truncated_workbook(manual_app, monkeypatch, boundary):
    import server.db

    _, client, headers, _, _ = manual_app
    consumed = []

    def oversized_source(*args, **kwargs):
        if boundary == "characters":
            consumed.append(1)
            yield {"occurred_at": NOW, "direction": "in", "text": "x" * 20_000_001}
        else:
            for index in range(100_002):
                consumed.append(index)
                yield {"occurred_at": NOW, "direction": "in", "text": "synthetic-row"}

    # Source size injection isolates the HTTP/XLSX limit, while the adjacent
    # 1001-row test and explicit load gate observe real database chunking.
    monkeypatch.setattr(server.db, "stream_rows", oversized_source)
    response = client.get("/api/monitor/export", headers=headers["a"], params={"start": "2026-10-13", "end": "2026-10-14"})
    assert response.status_code == 400
    assert "spreadsheet" not in response.headers.get("content-type", "")
    assert len(consumed) == (100_001 if boundary == "rows" else 1)


def test_non_string_manual_peer_is_a_client_error_without_state_change(manual_app):
    app, _, headers, _, _ = manual_app
    authorize(manual_app)
    before = app.state.inquiry.conversations("a")
    with TestClient(app, raise_server_exceptions=False) as client:
        response = client.post("/api/monitor/send", headers=headers["a"], json={"cp": [], "text": "合成消息", "request_id": "bad-peer"})
    assert response.status_code in (400, 422)
    assert app.state.inquiry.conversations("a") == before
