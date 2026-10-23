"""Developer regressions: retained facts and isolated retired execution scopes."""
import json

import pytest

from tests.institution_test_support import OWNER, selection
from tests.cutover.test_fresh_workspace import seed_command
from tests.cutover.test_im_catalog_selection import choose
from server.receiver_ledger import ReceiverLedger, account_lock
from relay.gateway_paths import CONTRACT, DIRECTORY_PATH, ORGANIZATION_PATH
from relay.protocol import digest


@pytest.mark.parametrize('state', ['QUEUED', 'PENDING', 'UNKNOWN', 'FAILED', 'CANCELLED', 'SENT'])
def test_switch_archives_projection_and_retains_original_facts(selection, state):
    s = selection
    gid, mid, tid = seed_command(s, state)
    store = s['repository']
    with store.db.transaction(write=False) as tx:
        message = dict(tx.execute('SELECT * FROM xj2_message WHERE thread_id=?', (tid,)).fetchone())
        command = dict(tx.execute("SELECT * FROM xj2_command WHERE command_id='old-command'").fetchone())
        audit = [dict(row) for row in tx.execute('SELECT * FROM xj2_audit').fetchall()]
    response = choose(s)
    assert response.status_code == 200, response.text
    with store.db.transaction(write=False) as tx:
        group = dict(tx.execute('SELECT * FROM xj2_group WHERE group_id=?', (gid,)).fetchone())
        peer = dict(tx.execute('SELECT * FROM xj2_peer WHERE member_id=?', (mid,)).fetchone())
        retained = dict(tx.execute("SELECT * FROM xj2_command WHERE command_id='old-command'").fetchone())
        assert group['archived'] == 1 and group['send_allowed'] == peer['send_allowed'] == 0
        assert dict(tx.execute('SELECT * FROM xj2_message WHERE thread_id=?', (tid,)).fetchone()) == message
        for old in audit:
            assert dict(tx.execute('SELECT * FROM xj2_audit WHERE audit_id=?', (old['audit_id'],)).fetchone()) == old
        for key in ('body_json', 'basis_json', 'content_hash', 'native_id', 'sent_at', 'created_at'):
            assert retained[key] == command[key]
        assert retained['state'] == ('CANCELLED' if state == 'QUEUED' else state)
        assert retained['cancel_requested'] == int(state in ('PENDING', 'UNKNOWN'))
    assert not store.binding_for(OWNER)['bot_enabled']
    assert s['api'].get('/api/monitor/counterparties').json()['data'] == []
    assert s['api'].get('/api/im/commands/unknown').json()['data'] == []


def test_retired_transport_keeps_unknown_download_and_upload_bytes(selection):
    s = selection
    old = s['runtime'].current_adapter(OWNER)
    with old.store.db() as db:
        seq = db.execute("INSERT INTO fg_download(scope,state,name,blob,code) VALUES(?,'UNKNOWN','lost',?,'unknown')",
                         (old.store.scope, b'original-unknown-input')).lastrowid
        db.execute("INSERT INTO fg_upload VALUES(?,?,'UNKNOWN','','outcome-unknown')", (old.store.scope, 'original-upload'))
        before = tuple(db.execute('SELECT * FROM fg_download WHERE seq=?', (seq,)).fetchone())
        upload = tuple(db.execute('SELECT * FROM fg_upload WHERE id=?', ('original-upload',)).fetchone())
    assert choose(s).status_code == 200
    with old.store.db() as db:
        assert tuple(db.execute('SELECT * FROM fg_download WHERE seq=?', (seq,)).fetchone()) == before
        assert tuple(db.execute('SELECT * FROM fg_upload WHERE id=?', ('original-upload',)).fetchone()) == upload
    assert not old.context_ready
    assert old not in s['app'].state.transport.adapters.values()


def test_automatic_native_slot_change_archives_without_retarget_or_resend(selection, monkeypatch):
    s = selection
    gid, mid, tid = seed_command(s)
    old = s['runtime'].current_adapter(OWNER)
    original = s['runtime'].directory.catalog()
    replacement = json.loads(json.dumps(original))
    account = next(a for w in replacement['workbenches'] if w['workbench_id'] == 'wb-a'
                   for a in w['accounts'] if a['account_uid'] == 'account-a')
    account['account_uid'] = 'new-native-login'
    monkeypatch.setattr(s['runtime'].directory, 'catalog', lambda: replacement)
    calls = list(s['wire'].calls)
    s['runtime']._reset_changed_native_accounts()
    assert s['wire'].calls == calls
    assert s['repository'].binding_for(OWNER) is None
    with s['repository'].db.transaction(write=False) as tx:
        assert tx.execute('SELECT archived FROM xj2_group WHERE group_id=?', (gid,)).fetchone()['archived'] == 1
        assert tx.execute("SELECT state FROM xj2_command WHERE command_id='old-command'").fetchone()['state'] == 'UNKNOWN'
        assert tx.execute('SELECT account_uid FROM xj2_thread WHERE thread_id=?', (tid,)).fetchone()['account_uid'] == 'account-a'
    assert not old.context_ready and old not in s['app'].state.transport.adapters.values()


def test_qidian_selector_filters_wechat_and_rejects_wechat_selection(selection, monkeypatch):
    s = selection
    directory = s['runtime'].directory
    original = directory.catalog()
    mixed = json.loads(json.dumps(original))
    mixed['workbenches'][0]['accounts'].append({
        'account_uid': 'wechat-only', 'account_id': 'wechat-slot', 'channel': 'wechat',
        'available': True, 'state': 'online', 'updated_at': s['clock'][0],
        'capabilities': ['negotiation-v22-watermark']})
    monkeypatch.setattr(directory, 'catalog', lambda: mixed)
    catalog = s['api'].get('/api/im/catalog').json()['data']
    assert all(a['channel'] == 'qidian' for w in catalog['workbenches'] for a in w['accounts'])
    before = s['repository'].binding_for(OWNER)
    response = choose(s, 'wb-a', 'wechat-only')
    assert response.status_code == 400
    assert s['repository'].binding_for(OWNER) == before


def test_reselected_same_account_excludes_retired_unknown_and_old_messages(selection):
    s = selection
    gid, mid, tid = seed_command(s)
    assert choose(s).status_code == 200
    assert choose(s, 'wb-a', 'account-a').status_code == 200
    assert s['repository'].snapshot(OWNER)['groups'] == []
    assert s['repository'].snapshot(OWNER)['threads'] == []
    assert s['repository'].raw_page(OWNER, tid)['messages'] == []
    assert s['app'].state.inquiry.commands(OWNER) == []
    calls = list(s['wire'].calls)
    s['app'].state.dispatcher.run_once()
    assert s['wire'].calls == calls
    store = s['repository']
    from server.negotiation.prices import LimitPolicy
    new_group = store.create_group(OWNER, bond_code='NEW.IB', our_side='BUY', target_volume='1000',
        trade_day='2026-10-23', policy=LimitPolicy(mode='ABSOLUTE', value='2'), send_allowed=False)
    new_member, new_tid = store.add_member(OWNER, new_group, account_uid='account-a', peer_uid='old-peer')
    assert new_tid == tid  # Existing account/peer uniqueness is retained.
    assert store.claim_analysis(OWNER, tid, force=True) is None
    new_message = store.ingest(OWNER, tid, event_id='new-native', source_seq=2, text='new context', trade_day='2026-10-23')
    claim = store.claim_analysis(OWNER, tid, force=True)
    assert [item['message_id'] for item in claim['snapshot']['messages']] == [new_message]
    assert [item['member_id'] for item in claim['snapshot']['members']] == [new_member]


@pytest.mark.parametrize('reselect_same', [False, True])
def test_original_late_sent_receipt_is_retained_without_current_analysis(selection, reselect_same):
    s = selection
    gid, mid, tid = seed_command(s)
    old = s['runtime'].current_adapter(OWNER)
    with s['repository'].write(OWNER) as tx:
        tx.execute("UPDATE xj2_command SET body_json=?,basis_json=? WHERE command_id='old-command'", (
            json.dumps({'kind': 'INITIAL', 'text': 'immutable old send', 'scope_members': [mid]}),
            json.dumps({'trade_day': '2026-10-22'})))
    assert choose(s).status_code == 200
    if reselect_same:
        assert choose(s, 'wb-a', 'account-a').status_code == 200
    payload = {'command_id': 'old-command', 'content_hash': 'a' * 64,
        'peer_id': 'old-peer', 'account_uid': 'account-a', 'status': 'sent',
        'evidence': 'synthetic exact native receipt', 'native_id': 'old-native-send', 'executed_at': s['now']}
    envelope = {'kind': 'receipt', 'owner_user_id': OWNER, 'workbench_id': 'wb-a',
        'account_uid': 'account-a', **{key: payload[key] for key in ('command_id', 'content_hash', 'peer_id')}, 'payload': payload}
    assert s['app'].state.inbox.apply(old.binding, envelope) == 'SENT'
    with s['repository'].db.transaction(write=False) as tx:
        command = dict(tx.execute("SELECT * FROM xj2_command WHERE command_id='old-command'").fetchone())
        message = dict(tx.execute("SELECT * FROM xj2_message WHERE native_id='old-native-send'").fetchone())
        thread = dict(tx.execute('SELECT * FROM xj2_thread WHERE thread_id=?', (tid,)).fetchone())
        assert command['state'] == 'SENT' and command['native_id'] == 'old-native-send'
        assert message['historical'] == 1 and message['text'] == 'immutable old send'
        assert thread['pending_since'] is None
    assert s['repository'].snapshot(OWNER)['groups'] == []
    assert s['repository'].snapshot(OWNER)['threads'] == []
    assert s['repository'].claim_analysis(OWNER, tid, force=True) is None


def test_config_replacement_retains_unbound_history_and_advances_future_binding_version(selection):
    s = selection
    gid, mid, tid = seed_command(s)
    old = s['repository'].binding_for(OWNER)
    status = s['api'].get('/api/im/mq-config').json()['data']
    response = s['api'].put('/api/im/mq-config', json={'base_url': 'https://replacement.invalid',
        'config_yaml': '', 'expected_revision': status['revision']})
    assert response.status_code == 200, response.text
    assert s['repository'].binding_for(OWNER) is None
    assert s['repository'].snapshot(OWNER)['groups'] == []
    assert s['repository'].snapshot(OWNER)['threads'] == []
    assert s['app'].state.inquiry.commands(OWNER) == []
    response = choose(s, 'wb-a', 'account-a', version=0)
    assert response.status_code == 200, response.text
    assert s['repository'].binding_for(OWNER)['version'] > old['version']
    with s['repository'].db.transaction(write=False) as tx:
        assert tx.execute("SELECT state FROM xj2_command WHERE command_id='old-command'").fetchone()['state'] == 'UNKNOWN'


def test_receiver_context_retires_execution_but_preserves_unknown_body_and_audit(tmp_path):
    from tests.cutover.test_session_reset_protocol import B, command, NOW
    receiver = ReceiverLedger(tmp_path / 'receiver.db', B, clock=lambda: NOW)
    receiver.activate_context('old-context', NOW * 1000000)
    old = command(context_id='old-context')
    receiver.execute(old, lambda _: {'status': 'unknown', 'evidence': 'native timed out'})
    with receiver._db() as db:
        original = tuple(db.execute('SELECT * FROM rx_command').fetchone())
        audit = [tuple(row) for row in db.execute('SELECT * FROM rx_audit')]
    receiver.activate_context('new-context', NOW * 1000000 + 1)
    with receiver._db() as db:
        assert tuple(db.execute('SELECT * FROM rx_command').fetchone()) == original
        assert [tuple(row) for row in db.execute('SELECT * FROM rx_audit')] == audit
        assert db.execute('SELECT COUNT(*) FROM rx_tombstone').fetchone()[0] == 1
    calls = []
    assert receiver.execute(old, lambda _: calls.append('must not send'))['status'] == 'unknown'
    assert calls == [] and receiver.permission(old) is False


def test_shared_directory_has_one_short_locked_consumer_and_private_business_state(selection):
    s = selection
    config = json.loads(s['connection'].read_text())
    config['directory_state_dir'] = str(s['tmp'] / 'shared-directory')
    s['connection'].write_text(json.dumps(config))
    s['runtime'].load()
    from server.institution_runtime import InstitutionRuntime
    other = InstitutionRuntime(s['repository'], s['app'].state.transport, s['connection'], s['tmp'] / 'other-state')
    assert other.directory.path == s['runtime'].directory.path == s['tmp'] / 'shared-directory' / s['keys'].institution_id / 'directory.db'
    assert other.state != s['runtime'].state
    calls = list(s['wire'].calls)
    with account_lock(str(other.directory.path) + '.consume', other.keys.binding('directory', 'directory'), blocking=False):
        s['runtime'].tick()
    assert not any(path in (DIRECTORY_PATH, ORGANIZATION_PATH) for path, _ in s['wire'].calls[len(calls):])
    s['runtime'].tick()
    assert any(path == DIRECTORY_PATH for path, _ in s['wire'].calls[len(calls):])


def test_receiver_same_uid_qidian_and_wechat_have_separate_scopes_and_locks(tmp_path):
    from dataclasses import replace
    from tests.cutover.test_session_reset_protocol import B, command, NOW
    wechat = replace(B, channel='wechat')
    path = tmp_path / 'shared-receiver.db'
    qid = ReceiverLedger(path, B, clock=lambda: NOW)
    wx = ReceiverLedger(path, wechat, clock=lambda: NOW)
    assert qid.scope == 'xunjia:' + digest([B.workbench_id, B.account_uid])
    assert qid.scope != wx.scope
    with account_lock(path, B, blocking=False):
        with pytest.raises(BlockingIOError):
            with account_lock(path, B, blocking=False):
                pytest.fail('Same channel lock must be exclusive')
        with account_lock(path, wechat, blocking=False):
            pass
    qid.execute(command(), lambda _: {'status': 'unknown', 'evidence': 'qidian old unknown'})
    calls = []
    result = wx.execute({**command(), 'channel': 'wechat'},
        lambda _: (calls.append('wechat') or {'status': 'sent', 'evidence': 'wechat native exact', 'native_id': 'wx-native'}))
    assert result['status'] == 'sent' and calls == ['wechat']
    with qid._db() as db:
        assert db.execute('SELECT status FROM rx_command WHERE scope=?', (qid.scope,)).fetchone()['status'] == 'unknown'
        assert db.execute('SELECT status FROM rx_command WHERE scope=?', (wx.scope,)).fetchone()['status'] == 'sent'


@pytest.mark.parametrize('operation', ['direct', 'confirm', 'preview', 'manual'])
def test_retired_request_ids_never_restore_old_current_ui(selection, operation):
    s = selection
    service = s['app'].state.inquiry
    items = [{'code': 'SYN-RETENTION', 'our_side': 'BUY', 'peer_ids': ['retention-peer'],
        'volume_wan': '1000', 'opening_yield_pct': '3.20', 'limit_yield_pct': '3.12',
        'step_bp': '2', 'pricing_basis': 'absolute'}]
    if operation in ('direct', 'manual'):
        result = service.send_direct(OWNER, 'retired-direct-request', items)
        if operation == 'manual':service.manual_send(OWNER, 'retention-peer', 'original manual text', 'retired-manual')
    else:
        preview = service.preview(OWNER, 'retired-preview', items)
        result = service.confirm(OWNER, preview['batch_id'], preview['version'], 'retired-confirm')
    assert choose(s).status_code == 200
    assert choose(s, 'wb-a', 'account-a').status_code == 200
    with pytest.raises((ValueError, PermissionError)):
        if operation == 'direct':service.send_direct(OWNER, 'retired-direct-request', items)
        elif operation == 'manual':service.manual_send(OWNER, 'retention-peer', 'original manual text', 'retired-manual')
        elif operation == 'confirm':service.confirm(OWNER, preview['batch_id'], preview['version'], 'retired-confirm')
        else:service.preview(OWNER, 'retired-preview', items)
    assert s['repository'].snapshot(OWNER)['groups'] == []
    with s['repository'].db.transaction(write=False) as tx:
        assert tx.execute('SELECT COUNT(*) AS n FROM xj2_batch').fetchone()['n'] == 1
        assert tx.execute('SELECT COUNT(*) AS n FROM xj2_group').fetchone()['n'] == 1


def test_financial_quote_and_source_scope_are_retained_byte_for_byte(selection):
    from server.negotiation.facts import Quote
    from server.negotiation.prices import Price
    s = selection
    gid, mid, tid = seed_command(s)
    with s['repository'].write(OWNER) as tx:
        message = tx.execute('SELECT message_id FROM xj2_message WHERE thread_id=?', (tid,)).fetchone()['message_id']
        tx.execute("INSERT INTO xj2_scope VALUES(?,?,'MENTION')", (message, mid))
        s['repository']._save_quote(tx, mid, Quote('ofr', price=Price('NET_PRICE', '100.123456'),
            volume_wan='1000.000001', revision=1, evidence={'price': [message], 'volume_wan': [message]}), 'original-analysis')
    with s['repository'].db.transaction(write=False) as tx:
        quote = dict(tx.execute('SELECT * FROM xj2_quote WHERE member_id=?', (mid,)).fetchone())
        scope = dict(tx.execute('SELECT * FROM xj2_scope WHERE member_id=?', (mid,)).fetchone())
    assert choose(s).status_code == 200
    with s['repository'].db.transaction(write=False) as tx:
        assert dict(tx.execute('SELECT * FROM xj2_quote WHERE member_id=?', (mid,)).fetchone()) == quote
        assert dict(tx.execute('SELECT * FROM xj2_scope WHERE member_id=?', (mid,)).fetchone()) == scope


def test_new_context_keeps_old_reset_facts_without_running_old_reset(tmp_path):
    from tests.cutover.test_session_reset_protocol import B, command, request, NOW
    from relay.protocol import session_reset_hash
    receiver = ReceiverLedger(tmp_path / 'reset-context.db', B, clock=lambda: NOW)
    receiver.activate_context('old-context', NOW * 1000000)
    old = {**request(), 'context_id': 'old-context'}
    old['request_hash'] = session_reset_hash(old)
    assert receiver.session_reset(old)['phase'] == 'blocked'
    with receiver._db() as db:
        original = tuple(db.execute('SELECT * FROM rx_reset').fetchone())
    context = receiver.activate_context('new-context', NOW * 1000000 + 1)
    assert context['account_epoch'] == 1
    assert receiver.pending_session_resets() == []
    with receiver._db() as db:
        assert tuple(db.execute('SELECT * FROM rx_reset').fetchone()) == original
    outcome = receiver.execute(command(context_id='new-context', account_epoch=1),
        lambda _: {'status': 'sent', 'evidence': 'synthetic exact native result', 'native_id': 'new-native'})
    assert outcome['status'] == 'sent'
    current = {**request(reset_id='new-reset', expected=1, epoch=2), 'context_id': 'new-context'}
    current['request_hash'] = session_reset_hash(current)
    assert receiver.session_reset(current, native_quiescent=lambda: True)['phase'] == 'barrier_ready'
    assert [item['reset_id'] for item in receiver.pending_session_resets()] == ['new-reset']


def test_qidian_runtime_never_consumes_wechat_business_queues(selection):
    from relay.gateway_paths import account_paths
    from server.institution_transport import InstitutionTransport
    s = selection
    binding = s['keys'].binding('wx-workbench', 'wx-account', 'wx-owner', channel='wechat')
    paths = account_paths(binding.workbench_id, binding.account_uid, channel='wechat')
    adapter = InstitutionTransport(binding, s['wire'].client(paths['uplink_path']),
        s['wire'].client(paths['downlink_path']), s['tmp'] / 'wechat-service.db',
        role='service', institution_owner=s['keys'].owner_id, clock=lambda: s['clock'][0])
    s['app'].state.transport.adapters[('wx-owner', 'wx-account', 'oss')] = adapter
    calls = list(s['wire'].calls)
    s['runtime'].tick()
    s['app'].state.inbox.poll_once()
    assert not any(path in paths.values() for path, _ in s['wire'].calls[len(calls):])
    with pytest.raises(PermissionError, match='Qidian'):
        s['app'].state.inbox.apply(s['runtime'].current_adapter(OWNER).binding, {'channel': 'wechat'})


def test_missing_binding_never_exposes_old_projection_even_without_retirement_audit(selection):
    s = selection
    gid, mid, tid = seed_command(s)
    with s['repository'].write(OWNER) as tx:
        tx.execute('DELETE FROM xj2_binding WHERE owner_id=?', (OWNER,))
    assert s['repository'].snapshot(OWNER)['groups'] == []
    assert s['repository'].snapshot(OWNER)['threads'] == []
    assert s['app'].state.inquiry.commands(OWNER) == []
    with pytest.raises(ValueError):
        s['repository'].raw_page(OWNER, tid)


def test_same_uid_http_monitor_cursor_export_and_capacity_start_empty(selection):
    from io import BytesIO
    from openpyxl import load_workbook
    s = selection
    gid, mid, tid = seed_command(s)
    s['repository'].ingest(OWNER, tid, event_id='second-old', source_seq=2,
        text='second old private message', trade_day='2026-10-22')
    assert choose(s).status_code == 200
    assert choose(s, 'wb-a', 'account-a').status_code == 200
    assert s['api'].get('/api/monitor/counterparties').json()['data'] == []
    assert s['api'].get('/api/monitor/detail').json()['data'] == []
    assert s['api'].get('/api/negotiation/capacity').json()['data']['active'] == 0
    chat = s['api'].get('/api/monitor/chat?cp=old-peer&limit=1').json()['data']
    assert chat['messages'] == [] and not chat['hasMore'] and chat['nextCursor'] is None
    export = s['api'].get('/api/monitor/export?start=2000-01-01&end=2100-01-01')
    assert export.status_code == 200, export.text
    assert len(list(load_workbook(BytesIO(export.content), read_only=True).active.values)) == 1
    s['repository'].ingest(OWNER, tid, event_id='new-visible', source_seq=3,
        text='new current message', trade_day='2026-10-23')
    chat = s['api'].get('/api/monitor/chat?cp=old-peer&limit=1').json()['data']
    assert [item['text'] for item in chat['messages']] == ['new current message']
    assert not chat['hasMore'] and chat['nextCursor'] is None
    export = s['api'].get('/api/monitor/export?start=2000-01-01&end=2100-01-01')
    rows = list(load_workbook(BytesIO(export.content), read_only=True).active.values)
    assert len(rows) == 2 and rows[1][2] == 'new current message'
    for scope in ({'member_id': mid}, {'group_id': gid}):
        with pytest.raises(ValueError, match='read only'):
            s['repository'].set_sending(OWNER, True, **scope)

